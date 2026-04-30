import logging
from pathlib import Path
from typing import List

from openpyxl import load_workbook

from models import InputRow

log = logging.getLogger("lemit_bot.input_reader")


def get_input_file(folder: Path) -> Path:
    """Locate the xlsx file to process, with interactive selection if multiple."""
    if not folder.exists():
        raise FileNotFoundError(
            f"Pasta de entrada nao encontrada: {folder}\n"
            "Crie a pasta e coloque sua planilha .xlsx dentro."
        )

    files = sorted(folder.glob("*.xlsx"))
    if not files:
        raise FileNotFoundError(
            f"Nenhum arquivo .xlsx encontrado em: {folder}\n"
            "Coloque sua planilha Excel na pasta de entrada."
        )

    if len(files) == 1:
        log.info("Arquivo encontrado: %s", files[0].name)
        return files[0]

    log.info("%d arquivos xlsx encontrados, solicitando selecao", len(files))
    return _select_file_interactive(files)


def _select_file_interactive(files: List[Path]) -> Path:
    from rich.console import Console
    from rich.table import Table

    console = Console()
    table = Table(title="Planilhas disponíveis em 'entrada/'", show_header=True, header_style="bold cyan")
    table.add_column("#", style="cyan", width=4, justify="right")
    table.add_column("Arquivo", style="white")
    table.add_column("Tamanho", style="dim", justify="right")

    for i, f in enumerate(files, 1):
        size = f.stat().st_size
        size_str = f"{size / 1024:.1f} KB" if size < 1_000_000 else f"{size / 1_000_000:.1f} MB"
        table.add_row(str(i), f.name, size_str)

    console.print(table)

    while True:
        try:
            choice = console.input(
                f"[bold cyan]Selecione o número do arquivo [1-{len(files)}] "
                f"(Enter para o primeiro): [/bold cyan]"
            ).strip()
            if choice == "":
                selected = files[0]
            else:
                idx = int(choice) - 1
                if not (0 <= idx < len(files)):
                    console.print(f"[red]Digite um número entre 1 e {len(files)}.[/red]")
                    continue
                selected = files[idx]
            log.info("Arquivo selecionado: %s", selected.name)
            return selected
        except ValueError:
            console.print("[red]Entrada inválida. Digite um número.[/red]")
        except (EOFError, KeyboardInterrupt):
            raise SystemExit(0)


def read_rows(file_path: Path, cpf_column: str, name_column: str) -> List[InputRow]:
    """Read and validate input Excel. Returns list of InputRow (skips empty rows)."""
    log.info("Lendo planilha: %s", file_path.name)

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    # Read header (first row)
    raw_header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not raw_header:
        raise ValueError("Planilha vazia ou sem cabeçalho.")

    headers = [str(h).strip().lower() if h is not None else "" for h in raw_header]

    cpf_col = cpf_column.strip().lower()
    name_col = name_column.strip().lower()

    if cpf_col not in headers:
        raise ValueError(
            f"Coluna '{cpf_column}' não encontrada na planilha.\n"
            f"Colunas disponíveis: {', '.join(h for h in headers if h)}\n"
            "Verifique o campo 'excel.cpf_column' no config.json."
        )
    if name_col not in headers:
        raise ValueError(
            f"Coluna '{name_column}' não encontrada na planilha.\n"
            f"Colunas disponíveis: {', '.join(h for h in headers if h)}\n"
            "Verifique o campo 'excel.name_column' no config.json."
        )

    cpf_idx = headers.index(cpf_col)
    name_idx = headers.index(name_col)

    rows: List[InputRow] = []
    for row_num, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def _cell(idx: int) -> str:
            val = row_values[idx] if idx < len(row_values) else None
            return str(val).strip() if val is not None else ""

        cpf_val = _cell(cpf_idx)
        nome_val = _cell(name_idx)

        if not cpf_val and not nome_val:
            log.debug("Linha %d: CPF e nome vazios, pulando", row_num)
            continue

        raw_data = {headers[i]: row_values[i] for i in range(len(headers)) if i < len(row_values)}

        rows.append(InputRow(row_index=row_num, nome=nome_val, cpf=cpf_val, raw_data=raw_data))

    wb.close()
    log.info("%d linhas carregadas", len(rows))
    return rows
