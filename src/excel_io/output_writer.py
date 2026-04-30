import logging
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import QueryResult

log = logging.getLogger("lemit_bot.output_writer")

COLUMNS = [
    "nome", "cpf",
    "telefone_1", "telefone_2", "telefone_3",
    "email_1", "email_2",
    "consulta_usada", "status_consulta", "observacao", "data_consulta",
]

_COL_WIDTHS = [30, 16, 18, 18, 18, 32, 32, 14, 22, 40, 20]

_STATUS_FILL = {
    "encontrado_por_cpf":  "C6EFCE",
    "encontrado_por_nome": "FFEB9C",
    "nao_encontrado":      "FFCCCC",
    "erro":                "D9D9D9",
}


class OutputWriter:
    def __init__(self, output_dir: Path, whatsapp_template: str = "Olá {nome}, tudo bem?"):
        output_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y-%m-%d_%H-%M")
        self.file_path = output_dir / f"resultado_{ts}.xlsx"

        self._whatsapp_template = whatsapp_template
        self._wb = Workbook()
        self._ws = self._wb.active
        self._ws.title = "Resultados"
        self._write_header()
        self._wb.save(self.file_path)
        log.info("Planilha de saida criada: %s", self.file_path)

    def _write_header(self) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        center = Alignment(horizontal="center")

        for col_idx, col_name in enumerate(COLUMNS, 1):
            cell = self._ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        self._ws.freeze_panes = "A2"

        for col_idx, width in enumerate(_COL_WIDTHS, 1):
            col_letter = self._ws.cell(row=1, column=col_idx).column_letter
            self._ws.column_dimensions[col_letter].width = width

    def append_row(self, result: QueryResult) -> None:
        row_num = self._ws.max_row + 1
        values = [
            result.nome, result.cpf,
            result.telefone_1, result.telefone_2, result.telefone_3,
            result.email_1, result.email_2,
            result.consulta_usada, result.status_consulta,
            result.observacao, result.data_consulta,
        ]

        color = _STATUS_FILL.get(result.status_consulta, "FFFFFF")
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        # telefone_1 fica na coluna 3 (índice 2 em 0-based)
        tel1_col = COLUMNS.index("telefone_1") + 1

        for col_idx, value in enumerate(values, 1):
            cell = self._ws.cell(row=row_num, column=col_idx, value=value)
            cell.fill = fill
            if col_idx == tel1_col and value:
                wa_url = _whatsapp_url(value, result.nome, self._whatsapp_template)
                if wa_url:
                    cell.hyperlink = wa_url
                    cell.font = Font(color="0563C1", underline="single")

        self._wb.save(self.file_path)
        log.debug("Linha salva: %s | %s", result.status_consulta, result.nome[:30])

    def close(self) -> None:
        self._wb.save(self.file_path)
        log.info("Planilha de saida finalizada: %s", self.file_path)


def _whatsapp_url(phone: str, nome: str, template: str) -> str:
    """Retorna URL wa.me com mensagem do template, ou '' se número inválido."""
    digits = re.sub(r"\D", "", phone)
    if not digits or digits.startswith("0800"):
        return ""
    # Adiciona DDI 55 se necessário (números brasileiros: 10 ou 11 dígitos)
    if len(digits) in (10, 11):
        digits = "55" + digits
    primeiro_nome = nome.split()[0] if nome else ""
    msg = template.format(nome=primeiro_nome)
    return f"https://wa.me/{digits}?text={quote(msg)}"
